"""
Модуль для разделения словарного .docx файла на отдельные файлы по буквам.
"""
import os
import re
from docx import Document
from docx.shared import Pt, RGBColor


def is_section_header(text):
    """
    Проверяет, является ли текст заголовком раздела.
    Заголовок — одиночная заглавная буква бурятского алфавита или специальный символ (ѲѲ).
    Бурятский алфавит включает: А-Я, Ё, Ө, Ү, Һ.
    """
    text = text.strip()
    if text == 'ѲѲ':
        return True
    if len(text) != 1:
        return False
    return bool(re.match(r'^[А-ЯЁӨҮҺѲ]$', text))


def copy_run_formatting(source_run, target_run):
    """Копирует форматирование текстового блока (run)."""
    try:
        target_run.bold = source_run.bold
        target_run.italic = source_run.italic
        target_run.underline = source_run.underline

        if source_run.font:
            if source_run.font.name:
                target_run.font.name = source_run.font.name
            if source_run.font.size:
                target_run.font.size = source_run.font.size
            if source_run.font.color and source_run.font.color.rgb:
                target_run.font.color.rgb = source_run.font.color.rgb
            if hasattr(source_run.font, 'all_caps'):
                target_run.font.all_caps = source_run.font.all_caps
            if hasattr(source_run.font, 'small_caps'):
                target_run.font.small_caps = source_run.font.small_caps
    except Exception as e:
        pass


def copy_paragraph_formatting(source_para, target_para):
    """Копирует форматирование параграфа."""
    try:
        if source_para.paragraph_format:
            fmt = target_para.paragraph_format
            src_fmt = source_para.paragraph_format

            if src_fmt.alignment:
                fmt.alignment = src_fmt.alignment
            if src_fmt.first_line_indent:
                fmt.first_line_indent = src_fmt.first_line_indent
            if src_fmt.left_indent:
                fmt.left_indent = src_fmt.left_indent
            if src_fmt.right_indent:
                fmt.right_indent = src_fmt.right_indent
            if src_fmt.space_before:
                fmt.space_before = src_fmt.space_before
            if src_fmt.space_after:
                fmt.space_after = src_fmt.space_after
            if src_fmt.line_spacing:
                fmt.line_spacing = src_fmt.line_spacing
    except Exception as e:
        pass


def save_section_to_file(letter, content, output_dir):
    """Сохраняет раздел в отдельный .docx файл с сохранением форматирования."""
    os.makedirs(output_dir, exist_ok=True)
    filename = os.path.join(output_dir, f"{letter}.docx")

    doc = Document()
    for text, source_para in content:
        if not text or not text.strip():
            continue

        new_para = doc.add_paragraph()

        if source_para.style:
            new_para.style = source_para.style

        copy_paragraph_formatting(source_para, new_para)

        if source_para.runs:
            for run in source_para.runs:
                new_run = new_para.add_run(run.text)
                copy_run_formatting(run, new_run)
        else:
            new_para.add_run(text)

    doc.save(filename)
    return filename


def split_dictionary(input_file, output_dir, log=None):
    """
    Разделяет словарь на отдельные файлы по буквам.

    Args:
        input_file: путь к исходному .docx файлу
        output_dir: папка для сохранения результатов
        log: функция обратного вызова для лог-сообщений — log(message, level)
             level: 'info', 'success', 'error'
    Returns:
        dict с результатами: {'sections': int, 'details': [{letter, count, file}]}
    """
    if log is None:
        log = lambda msg, level='info': None

    if not os.path.exists(input_file):
        log(f"Файл не найден: {input_file}", 'error')
        return None

    if not input_file.endswith('.docx'):
        log("Файл должен быть в формате .docx", 'error')
        return None

    log(f"Чтение файла: {os.path.basename(input_file)}", 'info')

    try:
        doc = Document(input_file)
    except Exception as e:
        log(f"Ошибка при чтении файла: {str(e)}", 'error')
        return None

    # Разбиение на секции
    sections = {}
    current_letter = None
    current_content = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        if is_section_header(text):
            if current_letter and current_content:
                sections[current_letter] = current_content
                current_content = []
            current_letter = text.strip()
        elif current_letter:
            current_content.append((paragraph.text, paragraph))

    # Последний раздел
    if current_letter and current_content:
        sections[current_letter] = current_content

    log(f"Найдено секций: {len(sections)}", 'info')

    # Сохранение файлов
    results = {'sections': len(sections), 'details': []}

    for letter, content in sections.items():
        try:
            filepath = save_section_to_file(letter, content, output_dir)
            results['details'].append({
                'letter': letter,
                'count': len(content),
                'file': filepath
            })
            log(f"Буква {letter} — {len(content)} записей → {os.path.basename(filepath)}", 'success')
        except Exception as e:
            log(f"Ошибка при сохранении буквы {letter}: {str(e)}", 'error')

    log(f"Сплит завершён. Создано файлов: {len(results['details'])}", 'success')
    return results
